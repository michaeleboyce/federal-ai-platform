"""Build a 10-row OMB-shaped fixture XLSX for the loader tests.

Covers every match category. Run from repo root:
  python3 tests/fixtures/_build_fixture.py
"""
from pathlib import Path

import openpyxl

OUT = Path(__file__).parent / "omb_consolidated_sample.xlsx"

# These are the EXACT header strings the real OMB file uses (verified
# against data/raw/2025_individually_reported_AI_use_cases.xlsx). The loader
# uses position-locked column ordering, so the prefixes here only need to
# satisfy column_maps.OMB_CONSOLIDATED_COLUMNS' prefix sanity check.
HEADERS = [
    "Agency Abbreviation", "Agency Name", "Use Case ID", "Use Case Name",
    "Bureau/Component", "Email Address",
    "Should this AI use case be withheld from public reporting?",
    "Stage of Development", "Is the AI use case high-impact?",
    "Justification", "Use Case Topic Area", "AI Classification",
    "What problem is the AI intended to solve?",
    "What are the expected benefits and positive outcomes",
    "Describe the AI system's outputs.",
    "Date when AI use case became operational",
    "Was the system involved in this use case purchased",
    "Vendor(s) Name",
    "Does this AI use case have an associated Authorization to Operate (ATO)?",
    "System(s) Name",
    "Describe any data used to train",
    "If the data is required to be publicly disclosed",
    "Does this AI use case involve PII",
    "If publicly available, provide the link to the PIA",
    "Which, if any, demographic variables",
    "Does this project include custom-developed code?",
    "If the code is open source, provide the link",
    "Has pre-deployment testing been conducted",
    "Has an AI impact assessment been completed",
    "What are the potential impacts",
    "Has as independent review",
    "Is there a process to conduct ongoing monitoring",
    "Has the agency established sufficient and periodic",
    "Does this AI use case have an appropriate fail-safe",
    "Is there an established appeal process",
    "What steps has the agency taken to consult",
]
assert len(HEADERS) == 36, f"expected 36 headers, got {len(HEADERS)}"


def row(abbr, name_full, uid, name, bureau="Test Bureau", **overrides):
    """Build a single 36-cell row with sensible defaults; override by key."""
    base = [
        abbr, name_full, uid, name, bureau, "user@example.gov",
        "a) No",                  # is_withheld
        "b) Pilot",               # stage_of_development
        "c) Not high-impact",     # is_high_impact
    ] + [None] * 27
    idx_map = {
        "stage": 7,
        "high_impact": 8,
        "topic": 10,
        "classification": 11,
        "vendor": 17,
        "have_ato": 18,
        "has_pii": 22,
        "has_custom_code": 25,
    }
    for k, v in overrides.items():
        base[idx_map[k]] = v
    return base


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidated Inventory"
    ws.append([None] * 36)        # row 1 (blank, matches real file)
    ws.append(HEADERS)            # row 2 (headers)

    # 1. Exact match (DB will have this name + agency exactly).
    ws.append(row("DOJ", "Department of Justice",
                  "DOJ-0001", "Test Veritone System"))
    # 2. Fuzzy match (DB has slight rename, ratio ~0.90+).
    ws.append(row("DOJ", "Department of Justice",
                  "DOJ-0002", "Test Veritone Audio Sys"))
    # 3. Suggested rename (NSF-style acronym expansion, ratio ~0.43 — falls
    #    in the 0.40–0.85 band).
    ws.append(row("NSF", "National Science Foundation",
                  7, "Technology, Innovation and Partnerships (TIP) "
                     "Microsoft (MS) Copilot Pilot"))
    # 4. OMB-only at agency we already track.
    ws.append(row("DHS", "Department of Homeland Security",
                  "DHS-9999", "Brand New DHS Use Case"))
    # 5. OMB-only at net-new agency.
    ws.append(row("PBGC", "Pension Benefit Guaranty Corp",
                  "PBGC-01", "PBGC Synthetic Data"))
    # 6. Drift case — DB has different stage + is_high_impact.
    ws.append(row("DOE", "Department of Energy",
                  "DOE-555", "Drift Test Case",
                  stage="c) Deployed", high_impact="a) High-impact"))
    # 7+8. Verbatim-duplicate row pair in OMB (PBGC-style).
    ws.append(row("PBGC", "Pension Benefit Guaranty Corp",
                  "PBGC-14", "Legislative and Regulatory Analysis"))
    ws.append(row("PBGC", "Pension Benefit Guaranty Corp",
                  "PBGC-15", "Legislative and Regulatory Analysis"))
    # 9. Empty ID (ED-style — the OMB file leaves use_case_id_omb blank).
    ws.append(row("ED", "Department of Education",
                  None, "Aidan Chat-bot"))
    # 10. STATE → State remap (the OMB file uses the abbreviation 'STATE').
    ws.append(row("STATE", "Department of State",
                  None, "AI Input in Translation"))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} with {ws.max_row} rows total ({ws.max_row - 2} data rows).")


if __name__ == "__main__":
    main()
